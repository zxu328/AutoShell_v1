import streamlit as st
import pandas as pd
import os
from openpyxl import Workbook, load_workbook
from datetime import datetime

USERS_FILE = "Users.xlsx"
REPAIRS_FILE = "AutoShield_Repairs.xlsx"
UPLOAD_DIR = "uploads"  # folder to store user images

# =========================================================
# Initialize Users.xlsx (map first 3 customers + admin user)
# =========================================================
def init_user_file():
    if not os.path.exists(USERS_FILE):
        df = pd.read_excel(REPAIRS_FILE)
        rows = df.to_dict(orient="records")

        wb = Workbook()
        ws = wb.active
        ws.title = "Users"
        ws.append(["UserID", "Username", "Password", "CustomerName", "CustomerEmail", "Role"])

        default_passwords = ["pass1", "pass2", "pass3"]

        for i, row in enumerate(rows[:3]):
            uid = i + 1
            username = f"user{i+1}"
            ws.append([
                uid,
                username,
                default_passwords[i],
                row["CustomerName"],
                row["CustomerEmail"],
                "user"
            ])
        
        # Add admin user
        ws.append([4, "user4", "adminpass", "Admin", "admin@autos.com", "admin"])
        wb.save(USERS_FILE)

        # Create main uploads directory
        if not os.path.exists(UPLOAD_DIR):
            os.makedirs(UPLOAD_DIR)

# =========================================================
# User helpers
# =========================================================
def load_users_df():
    return pd.read_excel(USERS_FILE)

def check_login(username, password):
    df = load_users_df()
    match = df[(df["Username"] == username) & (df["Password"] == password)]
    if not match.empty:
        row = match.iloc[0]
        return True, {
            "CustomerName": row["CustomerName"],
            "CustomerEmail": row["CustomerEmail"],
            "Role": row["Role"]
        }
    return False, {}

# =========================================================
# Login Page
# =========================================================
def show_login_page():
    st.title("🔐 AutoShield Login")

    username = st.text_input("Username")
    password = st.text_input("Password", type="password")

    if st.button("Login"):
        ok, user_map = check_login(username, password)
        if ok:
            st.session_state["logged_in"] = True
            st.session_state["username"] = username
            st.session_state["cust_name"] = user_map["CustomerName"]
            st.session_state["cust_email"] = user_map["CustomerEmail"]
            st.session_state["role"] = user_map["Role"]
            st.success("Login successful!")
        else:
            st.error("Invalid username or password.")

# =========================================================
# User Dashboard (with image upload)
# =========================================================
def show_dashboard():
    st.title("🚗 AutoShield Repair Dashboard")

    username = st.session_state["username"]
    cust_name = st.session_state["cust_name"]
    cust_email = st.session_state["cust_email"]
    role = st.session_state["role"]

    st.write(f"Welcome **{username}** — Customer: **{cust_name}**")

    # Load repair jobs
    df = pd.read_excel(REPAIRS_FILE)
    filtered = df if role == "admin" else df[df["CustomerEmail"].str.lower() == cust_email.lower()]

    st.subheader("Repair Job(s)")
    st.dataframe(filtered)

    if filtered.empty:
        st.info("No repair jobs found for your account.")
        return

    # =========================================================
    # Add Message Section
    # =========================================================
    st.subheader("📨 Add a Message / Note to Repair Job")
    job_ids = list(filtered["JobID"])
    selected_job = st.selectbox("Select Job ID", job_ids)
    subject = st.text_input("Subject")
    body = st.text_area("Message Body")

    if st.button("Submit Message"):
        wb = load_workbook(REPAIRS_FILE)
        if "Messages" not in wb.sheetnames:
            ws = wb.create_sheet("Messages")
            ws.append(["MessageID", "JobID", "PostedBy", "PostedAt", "Subject", "Body"])
        else:
            ws = wb["Messages"]

        message_id = ws.max_row
        timestamp = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        ws.append([message_id, selected_job, cust_name, timestamp, subject, body])
        wb.save(REPAIRS_FILE)
        st.success("Message submitted successfully!")

    # =========================================================
    # Upload Images Section
    # =========================================================
    st.subheader("📷 Upload Images")
    user_folder = os.path.join(UPLOAD_DIR, username)
    os.makedirs(user_folder, exist_ok=True)  # create folder if missing

    uploaded_files = st.file_uploader("Upload image(s)", type=["png", "jpg", "jpeg"], accept_multiple_files=True)
    if uploaded_files:
        for file in uploaded_files:
            file_path = os.path.join(user_folder, file.name)
            with open(file_path, "wb") as f:
                f.write(file.getbuffer())
        st.success(f"{len(uploaded_files)} image(s) uploaded successfully!")

    # Show uploaded images
    st.subheader("🖼️ Your Uploaded Images")
    user_images = os.listdir(user_folder)
    if user_images:
        for img_name in user_images:
            st.image(os.path.join(user_folder, img_name), caption=img_name)
    else:
        st.info("No images uploaded yet.")

    # =========================================================
    # Message History
    # =========================================================
    st.subheader("📄 Message History")
    wb = load_workbook(REPAIRS_FILE)
    if "Messages" in wb.sheetnames:
        messages = pd.DataFrame(wb["Messages"].values)
        messages.columns = messages.iloc[0]
        messages = messages[1:]
        if role != "admin":
            messages = messages[messages["JobID"].isin(filtered["JobID"].tolist())]
        messages["PostedAt"] = pd.to_datetime(messages["PostedAt"])
        messages = messages.sort_values("PostedAt", ascending=False)
        st.dataframe(messages)

    st.markdown("---")
    col1, col2 = st.columns(2)
    with col1:
        if st.button("Logout"):
            st.session_state.clear()
            st.success("Logged out.")
    with col2:
        st.info(f"Linked Customer Email: {cust_email}")

# =========================================================
# MAIN
# =========================================================
def main():
    st.set_page_config(page_title="AutoShield System", layout="centered")
    init_user_file()

    if "logged_in" not in st.session_state:
        st.session_state["logged_in"] = False

    if st.session_state["logged_in"]:
        show_dashboard()
    else:
        show_login_page()

if __name__ == "__main__":
    main()
